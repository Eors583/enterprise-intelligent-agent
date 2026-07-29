from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import RGBColor


SOURCE = Path(os.environ["DOCX_SOURCE"])
OUTPUT = Path(os.environ["DOCX_OUTPUT"])


PARAGRAPH_REPLACEMENTS: dict[int, str] = {
    0: "蓝血企业协同智能系统（BMS-AI）技术需求与实现方案",
    1: "全栈自研版｜自研智能体平台 + 自研企业知识库 + 模型 API 接入",
    2: "技术需求规格、总体架构、详细设计、治理体系与实施路线",
    3: (
        "方案定位：智能体运行、知识库、RAG、工具网关、权限治理、评测与运营平台全部自主研发；"
        "基础大模型不自研，通过统一模型 API 网关按需接入，避免依赖任何厂商的智能体或知识库平台。"
    ),
    33: "系统覆盖范围如下。",
    73: "员工登录Web、桌面端、移动端或已接入的协同渠道，身份服务解析员工、组织、岗位及当前有效角色任命。",
    85: "4.1 核心需求",
    124: "5.1 底座需求",
    207: "8.1 自研总体技术路线",
    208: (
        "本方案采用“应用与数据能力自研、基础模型通过API接入”的技术路线。企业自主建设智能体控制平面、"
        "Agent Runtime、任务编排、知识采集与RAG、工具网关、权限策略、审计、评测和运营体系；"
        "大模型通过统一模型API网关接入，不采购或依赖任何厂商的智能体开发平台、托管知识库或托管RAG。"
    ),
    209: "8.1.1 自研角色智能体平台",
    210: (
        "Agent Runtime：自主实现会话上下文、规划执行、状态机、任务队列、步数/时间/预算限制、"
        "失败重试、人工接管和多智能体协同协议；核心运行状态保存在企业自有数据库。"
    ),
    211: (
        "配置中心：自主实现角色模板、System Prompt、模型参数、知识域、工具策略、输出Schema和审批规则的"
        "版本化管理，支持草稿、测试、发布、灰度、回滚与审计。"
    ),
    212: (
        "自研知识库与RAG：自主实现文档接入、对象存储、解析、OCR、切分、Embedding、关键词/向量混合检索、"
        "重排、权限前置过滤、证据组装、引用校验、版本管理和质量评测。"
    ),
    213: (
        "工具调用：自主建设工具注册中心和工具网关，对内部IT系统API执行身份传递、参数校验、最小权限、"
        "幂等、限流、DryRun、人工审批、执行回执和补偿控制。"
    ),
    214: (
        "渠道适配：通过自研开放接口和适配器接入Web、桌面端、移动端及现有协同软件；渠道仅承担身份、消息和"
        "入口，不承载智能体编排、知识检索或企业权限判断。"
    ),
    217: (
        "自研边界：企业负责智能体平台、知识库、RAG、权限、流程、工具、审计、评测及数据资产；"
        "模型服务商只通过标准API提供推理能力。所有模型请求必须经过企业模型网关，业务系统不得直接调用模型。"
    ),
    227: "8.6 协同渠道与工作台集成",
    228: (
        "采用统一企业智能工作台，根据登录员工和有效任命路由到不同角色智能体；不为每个角色建设独立机器人，"
        "避免身份、权限、知识和运营配置分散。"
    ),
    229: (
        "渠道侧只负责身份登录、消息触达、待办通知和工作台入口；复杂目标、任务、纠偏、知识引用、审批和能力界面"
        "由自研Web/桌面工作台提供。"
    ),
    230: (
        "渠道适配层统一处理消息模板、频率控制、免打扰、签名验真、失败重试、回执和幂等；"
        "新增渠道只实现适配器，不改变Agent Runtime和知识服务。"
    ),
    231: (
        "外部渠道用户标识与企业EmployeeId建立受控映射；员工离职、组织变更或授权失效事件必须触发会话失效、"
        "权限回收和缓存清理。"
    ),
    276: "集成测试：HRIS、CRM、ERP、BPM、协同渠道、自研知识服务和模型API服务。",
    285: "12.1 总体实施路径",
    294: "通过标准适配器接入企业现有协同渠道，形成统一、可替换的智能办公入口。",
    326: "14.5 参考技术标准与内部规范",
    327: "模型API规范：企业模型网关统一封装Chat、Embedding、Rerank、流式输出、工具调用和错误码。",
    328: "知识服务规范：统一文档、版本、切片、索引、权限、检索、引用和反馈接口，不对业务暴露底层向量库。",
    329: "Agent协议规范：统一AgentRun、上下文、计划步骤、工具调用、审批、事件和审计数据结构。",
    330: "对象存储与检索规范：采用企业可控的S3兼容对象接口、关系数据库、全文索引和向量索引。",
    331: "说明：模型供应商可替换；接入前必须完成API兼容性、数据保留、SLA、成本、限流和退出能力评估。",
}


TABLE_REPLACEMENTS: dict[tuple[int, int, int], str] = {
    (0, 1, 1): "V3.0 全栈自研版",
    (0, 2, 1): "技术需求与总体实现方案（自研建设基线）",
    (0, 5, 1): "2026-07-28",
    (5, 1, 1): "Web工作台、桌面端、移动端、协同渠道适配器、开放API",
    (5, 1, 2): "身份登录、消息触达、工作台嵌入、第三方调用；渠道能力与核心平台解耦。",
    (21, 0, 0): "自研组件",
    (21, 0, 1): "主要职责",
    (21, 0, 2): "关键实现",
    (21, 0, 3): "控制边界",
    (21, 1, 0): "智能体控制平面",
    (21, 1, 1): "角色模板、配置、发布、路由、策略、运营",
    (21, 1, 2): "版本化配置、灰度、回滚、租户隔离、评测门禁",
    (21, 1, 3): "不在Prompt中固化业务主数据和权限",
    (21, 2, 0): "Agent Runtime",
    (21, 2, 1): "上下文、状态机、任务编排、多智能体协同、人工接管",
    (21, 2, 2): "队列驱动、幂等、超时、重试、预算与步数限制",
    (21, 2, 3): "模型只负责推理，运行状态由企业平台掌控",
    (21, 3, 0): "自研知识与RAG平台",
    (21, 3, 1): "采集、解析、切分、索引、检索、重排、引用、治理",
    (21, 3, 2): "对象存储＋关系库＋全文索引＋向量索引＋检索评测",
    (21, 3, 3): "权限在召回前过滤；向量库不作为权限和主数据权威",
}


GLOBAL_REPLACEMENTS = {
    "修订增强版": "全栈自研版",
    "企业微信或Web工作台": "Web、桌面端、移动端或协同渠道工作台",
    "企业微信等协同生态": "企业现有协同渠道",
    "企微消息": "协同渠道消息",
    "企微UserId": "外部渠道UserId",
    "企微接口": "渠道接口",
    "企微侧": "渠道侧",
    "企微入口": "统一工作台入口",
    "腾讯云": "",
    "腾讯": "",
    "混元": "外部大模型",
}


FORBIDDEN = ("腾讯", "混元", "ADP", "LangChain", "AutoGen", "微搭", "企业微信", "企微")


def replace_paragraph_text(paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.text = text
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.highlight_color = None


def replace_cell_text(cell, text: str) -> None:
    p = cell.paragraphs[0]
    replace_paragraph_text(p, text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def normalize_runs(document: Document) -> None:
    containers = [document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                containers.append(cell.paragraphs)
    for section in document.sections:
        containers.extend([section.header.paragraphs, section.footer.paragraphs])
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    containers.append(cell.paragraphs)
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    containers.append(cell.paragraphs)

    for paragraphs in containers:
        for paragraph in paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.highlight_color = None


def apply_global_replacements_to_paragraph(paragraph) -> None:
    text = paragraph.text
    updated = text
    for old, new in GLOBAL_REPLACEMENTS.items():
        updated = updated.replace(old, new)
    if updated != text:
        replace_paragraph_text(paragraph, updated)


def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs


def main() -> None:
    document = Document(SOURCE)

    for index, value in PARAGRAPH_REPLACEMENTS.items():
        replace_paragraph_text(document.paragraphs[index], value)

    for (table_index, row_index, column_index), value in TABLE_REPLACEMENTS.items():
        replace_cell_text(document.tables[table_index].cell(row_index, column_index), value)

    for paragraph in all_paragraphs(document):
        apply_global_replacements_to_paragraph(paragraph)

    normalize_runs(document)

    document.core_properties.title = "蓝血企业协同智能系统（BMS-AI）技术需求与实现方案（全栈自研版）"
    document.core_properties.subject = "自研智能体平台、自研企业知识库与RAG、模型API接入"
    document.core_properties.comments = (
        "移除托管智能体和托管知识库平台依赖；智能体、知识、RAG、工具、治理与评测全部自研。"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)

    reopened = Document(OUTPUT)
    remaining = []
    for paragraph in all_paragraphs(reopened):
        for token in FORBIDDEN:
            if token in paragraph.text:
                remaining.append((token, paragraph.text))
    if remaining:
        details = "\n".join(f"{token}: {text}" for token, text in remaining)
        raise RuntimeError(f"Forbidden vendor/platform references remain:\n{details}")

    print(f"Created: {OUTPUT}")
    print(f"Paragraphs: {len(reopened.paragraphs)}; tables: {len(reopened.tables)}")


if __name__ == "__main__":
    main()
