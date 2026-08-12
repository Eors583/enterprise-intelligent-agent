from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / ".data" / "knowledge-smoke" / "fixtures"
OUTPUT.mkdir(parents=True, exist_ok=True)


def build_image() -> Path:
    path = OUTPUT / "incident-dashboard.png"
    image = Image.new("RGB", (1200, 500), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, 1170, 470), outline="navy", width=5)
    draw.text((70, 70), "Enterprise Incident Dashboard", fill="navy")
    draw.text((70, 150), "Owner: Platform Operations", fill="black")
    draw.text((70, 220), "P1 response target: 15 minutes", fill="black")
    draw.text((70, 290), "Escalation: Security Director -> COO", fill="black")
    draw.text((70, 360), "Evidence ID: INC-POLICY-2026", fill="black")
    image.save(path)
    return path


def build_pdf(image_path: Path) -> None:
    path = OUTPUT / "incident-response-policy.pdf"
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm
    )
    story = [
        Paragraph("Enterprise Incident Response Policy 2026", styles["Title"]),
        Paragraph("Policy owner: Platform Operations", styles["Heading2"]),
        Paragraph(
            "A priority-one incident must be acknowledged within 15 minutes. "
            "The incident commander is the Platform Operations lead.",
            styles["BodyText"],
        ),
        Spacer(1, 8 * mm),
        Table(
            [
                ["Severity", "Response target", "Escalation"],
                ["P1", "15 minutes", "Security Director and COO"],
                ["P2", "60 minutes", "Department head"],
                ["P3", "1 business day", "Service owner"],
            ],
            colWidths=[35 * mm, 45 * mm, 80 * mm],
            style=TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightblue),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            ),
        ),
        PageBreak(),
        Paragraph("Escalation evidence", styles["Heading1"]),
        PdfImage(str(image_path), width=165 * mm, height=68.75 * mm),
        Spacer(1, 8 * mm),
        Paragraph(
            "For P1 incidents, the Security Director reviews containment evidence. "
            "The COO approves customer-facing recovery statements.",
            styles["BodyText"],
        ),
    ]
    document.build(story)


def build_docx() -> None:
    path = OUTPUT / "service-ownership-handbook.docx"
    document = Document()
    document.add_heading("Service Ownership Handbook", 0)
    document.add_heading("Payment Service", level=1)
    document.add_paragraph("Owner: Finance Platform Team")
    document.add_paragraph("Backup owner: Platform Operations")
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Service", "Owner", "SLA"]):
        cell.text = value
    for values in [
        ("Payment API", "Finance Platform Team", "99.95%"),
        ("Identity API", "Identity Team", "99.90%"),
    ]:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    document.add_heading("Dependency", level=1)
    document.add_paragraph("Payment API depends on Identity API for operator authentication.")
    document.save(path)


def build_xlsx() -> None:
    path = OUTPUT / "quarterly-service-metrics.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Service Metrics"
    sheet.append(["Service", "Q1 uptime", "Q2 uptime", "Average"])
    sheet.append(["Payment API", 0.999, 0.998, "=AVERAGE(B2:C2)"])
    sheet.append(["Identity API", 0.997, 0.999, "=AVERAGE(B3:C3)"])
    sheet.append(["Knowledge API", 0.995, 0.998, "=AVERAGE(B4:C4)"])
    owners = workbook.create_sheet("Owners")
    owners.append(["Service", "Department", "Primary owner"])
    owners.append(["Payment API", "Finance Platform", "Lin Xiao"])
    owners.append(["Identity API", "Enterprise Security", "Zhou Rui"])
    owners.append(["Knowledge API", "AI Platform", "Chen Yao"])
    workbook.save(path)


def build_pptx() -> None:
    path = OUTPUT / "knowledge-rollout-plan.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Enterprise Knowledge Rollout"
    slide.placeholders[1].text = (
        "Phase 1: ingestion and citations\n"
        "Phase 2: hybrid retrieval and rerank\n"
        "Phase 3: evaluation and continuous improvement"
    )
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Acceptance targets"
    table = slide.shapes.add_table(4, 2, Inches(1), Inches(1.8), Inches(8), Inches(3)).table
    values = [
        ("Metric", "Target"),
        ("Citation accuracy", ">= 95%"),
        ("Recall@10", ">= 85%"),
        ("P95 retrieval latency", "< 1500 ms"),
    ]
    for row_index, row in enumerate(values):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    presentation.save(path)


if __name__ == "__main__":
    dashboard = build_image()
    build_pdf(dashboard)
    build_docx()
    build_xlsx()
    build_pptx()
    for fixture in sorted(OUTPUT.iterdir()):
        print(f"{fixture.name}\t{fixture.stat().st_size}")
